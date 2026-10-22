"""Fixture tests for build-ema-data.py — EMA report + EU Union Register parsing.

The script is a command-line tool that runs at import time, so it is copied
into a temporary folder (together with agephrases.py, which it imports) and
run there as a subprocess: that keeps its side files (.sources.json, the
eu-register-decisions.json cache) out of the real repo and keeps the test
independent of the real smpc-data/ folder.

Covered: the positional column layout (header check), authorised / opinion /
negative / withdrawn rows, the Union Register fast path (a positive opinion the
Commission has already decided is published as authorised with ec="register"),
the register's last-good-download cache used during a Commission outage, and
the outage page itself ("Server temporarily unavailable" served with HTTP 200).
"""
import datetime as dt
import json
import os
import shutil
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _helpers import SCRIPTS_DIR, run_script  # noqa: E402

HEADERS = {
    0: "Category", 1: "Name of medicine", 2: "EMA product number", 3: "Medicine status", 4: "Opinion status",
    6: "International non-proprietary name (INN) / common name", 7: "Active substance",
    8: "Therapeutic area (MeSH)", 11: "ATC code (human)", 15: "Therapeutic indication",
    16: "Accelerated assessment", 18: "Advanced therapy", 19: "Biosimilar", 20: "Conditional approval",
    21: "Exceptional circumstances", 22: "Generic", 23: "Orphan medicine", 24: "PRIME: priority medicine",
    25: "Marketing authorisation developer / applicant / holder", 26: "European Commission decision date",
    29: "Opinion adopted date", 30: "Withdrawal of application date", 31: "Marketing authorisation date",
    32: "Refusal of marketing authorisation date",
    33: "Withdrawal / expiry / revocation / lapse of marketing authorisation date",
    34: "Suspension of marketing authorisation date", 38: "Medicine URL",
}
NCOL = 40
EPAR = "https://www.ema.europa.eu/en/medicines/human/EPAR/"


def row(cells):
    """A sheet row: {column index: value} -> list of NCOL cells."""
    r = [None] * NCOL
    for idx, v in cells.items():
        r[int(idx)] = v
    return r


def medicine(name, num, status, *, opinion=None, op_date=None, ma_date=None, ec_date=None, inn=None,
             withdrawn=None, extra=None):
    cells = {0: "Human", 1: name, 2: f"{num}/0000" if status == "Opinion" else num, 3: status, 4: opinion,
             6: inn or (name.lower() + "mab"), 7: inn or (name.lower() + "mab"), 8: "Test area", 11: "L01XX99",
             15: f"{name} is indicated for the treatment of adults with test disease.", 25: f"{name} Holder B.V.",
             26: ec_date, 29: op_date, 31: ma_date, 33: withdrawn, 38: EPAR + name.lower()}
    for k in list(cells):
        if cells[k] is None:
            del cells[k]
    if extra:
        cells.update(extra)
    return row(cells)


def write_report(path, medicines, generated=dt.datetime(2026, 10, 20)):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Medicine"
    ws.append(row({0: "Human medicines", 3: generated}))          # row 1: title + generation date
    for _ in range(7):                                             # rows 2-8: EMA's preamble
        ws.append(row({0: "…"}))
    ws.append(row(HEADERS))                                      # row 9: header (index 8)
    ws.append(row({0: "Human medicines only"}))                  # row 10 (index 9): skipped by the script
    for m in medicines:
        ws.append(m)
    wb.save(path)


def register_record(ema_num, date, kind="Centralised - Authorisation", cat="ch", page_id="2049"):
    return json.dumps({"cat": cat, "name": {"id": page_id, "en": "Test"},
                       "decision": {"ema_num": ema_num, "type": kind, "date": date}})


def write_register(path, records, wrap=True):
    lines = ['<html><body><script>var decisions = [']
    lines += [r + "," for r in records]
    lines += ['];</script></body></html>']
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines if wrap else records) + "\n")


OUTAGE_PAGE = "<html><head><title>Server temporarily unavailable</title></head><body>Sorry.</body></html>\n"


class BuildEmaDataTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = self.tmp.name
        # a private copy of the script + its helper module, so .sources.json and
        # the register cache land in the temp folder, not in the repo
        self.script = os.path.join(self.dir, "build-ema-data.py")
        shutil.copy(os.path.join(SCRIPTS_DIR, "build-ema-data.py"), self.script)
        shutil.copy(os.path.join(SCRIPTS_DIR, "agephrases.py"), os.path.join(self.dir, "agephrases.py"))
        self.xlsx = os.path.join(self.dir, "ema.xlsx")
        self.out = os.path.join(self.dir, "out", "ema-medicines.json")
        os.makedirs(os.path.dirname(self.out))
        self.cache = os.path.join(os.path.dirname(self.out), "eu-register-decisions.json")
        self.sources = os.path.join(self.dir, ".sources.json")
        write_report(self.xlsx, [
            medicine("Alphazon", "EMEA/H/C/006001", "Authorised", ma_date=dt.datetime(2026, 9, 1), op_date=dt.datetime(2026, 6, 25)),
            # authorised, MA-date column empty: the Commission decision date is the authorisation date
            medicine("Ecfallback", "EMEA/H/C/006005", "Authorised", ec_date=dt.datetime(2026, 7, 7), op_date=dt.datetime(2026, 4, 30)),
            # positive opinion the Commission has already decided (Union Register)
            medicine("Betarix", "EMEA/H/C/006002", "Opinion", opinion="Positive", op_date=dt.datetime(2026, 8, 20)),
            # positive opinion, decision still pending
            medicine("Gammaflu", "EMEA/H/C/006003", "Opinion", opinion="Positive", op_date=dt.datetime(2026, 9, 18)),
            # positive opinion whose only register decision predates the opinion (an older procedure) -> still pending
            medicine("Deltavir", "EMEA/H/C/006004", "Opinion", opinion="Positive", op_date=dt.datetime(2026, 9, 25)),
            medicine("Negatol", "EMEA/H/C/006006", "Opinion", opinion="Negative", op_date=dt.datetime(2026, 9, 18)),
            medicine("Oldone", "EMEA/H/C/001000", "Withdrawn", ma_date=dt.datetime(2015, 1, 1), withdrawn=dt.datetime(2026, 3, 3)),
            # veterinary rows are ignored
            row({0: "Veterinary", 1: "Dogmed", 2: "EMEA/V/C/000001", 3: "Authorised", 31: dt.datetime(2026, 9, 1), 38: EPAR + "dogmed"}),
        ])

    def tearDown(self):
        self.tmp.cleanup()

    def build(self, register_path):
        return run_script(self.script, self.xlsx, self.out, register_path or "", cwd=self.dir)

    def read(self):
        with open(self.out, encoding="utf-8") as f, open(self.sources, encoding="utf-8") as g:
            return json.load(f), json.load(g)

    def cache_decisions(self):
        with open(self.cache, encoding="utf-8") as f:
            return json.load(f)["decisions"]

    # ---- EMA report ---------------------------------------------------------------
    def test_report_rows_are_classified(self):
        self.build(None)
        data, sources = self.read()
        self.assertEqual(data["generated"], "2026-10-20")
        names = lambda key: [m["n"] for m in data[key]]  # noqa: E731
        self.assertEqual(names("authorised"), ["Alphazon", "Ecfallback"])
        self.assertEqual(names("pipeline"), ["Deltavir", "Gammaflu", "Betarix"])      # newest opinion first
        self.assertEqual(names("gone"), ["Negatol", "Oldone"])                   # newest event first
        alpha = data["authorised"][0]
        self.assertEqual((alpha["d"], alpha["op"], alpha["num"]), ("2026-09-01", "2026-06-25", "EMEA/H/C/006001"))
        self.assertEqual(data["authorised"][1]["d"], "2026-07-07")                   # EC decision date fallback
        self.assertEqual(data["pipeline"][2]["num"], "EMEA/H/C/006002")               # suffix /0000 stripped
        self.assertTrue(all(m["outcome"] == "positive" for m in data["pipeline"]))
        self.assertEqual((data["gone"][0]["st"], data["gone"][0]["e"]), ("Negative CHMP opinion", "2026-09-18"))
        self.assertEqual((data["gone"][1]["st"], data["gone"][1]["e"]), ("Withdrawn", "2026-03-03"))
        self.assertIn("alphazonmab", data["byInn"])
        self.assertEqual(sources["euRegisterDecisions"], 0)
        self.assertFalse(sources["euRegisterFromCache"])
        self.assertFalse(os.path.exists(self.cache))

    def test_changed_column_layout_refuses_to_build(self):
        wb = openpyxl.load_workbook(self.xlsx)
        ws = wb["Medicine"]
        ws.cell(row=9, column=32).value = "Some new column"          # the MA-date header moved
        wb.save(self.xlsx)
        res = run_script(self.script, self.xlsx, self.out, "", cwd=self.dir, check=False)
        self.assertNotEqual(res.returncode, 0)
        self.assertIn("layout changed", res.stderr)
        self.assertFalse(os.path.exists(self.out))

    # ---- EU Union Register ---------------------------------------------------------
    def test_register_decision_publishes_pending_medicine_as_authorised(self):
        reg = os.path.join(self.dir, "reg_last.htm")
        write_register(reg, [
            register_record("EMEA/H/C/006002/0000", "01/10/2026"),                         # Betarix: decided
            register_record("EMEA/H/C/006004/0000", "01/05/2026"),                         # Deltavir: older than its opinion
            register_record("EMEA/H/C/006001/II/0007", "15/09/2026", kind="Centralised - Variation"),  # not an authorisation
            register_record("EMEA/V/C/000001", "02/10/2026", cat="v"),                      # veterinary
            '{"cat":"ch","name":{"id":"1"},"decision":{"type":"Centralised - Authorisation","date":"03/10/2026"}}',  # no EMA number
            '{"cat":"ch", this is not json',                                                 # must be skipped, not fatal
        ])
        res = self.build(reg)
        data, sources = self.read()
        self.assertEqual([m["n"] for m in data["pipeline"]], ["Deltavir", "Gammaflu"])
        beta = next(m for m in data["authorised"] if m["n"] == "Betarix")
        self.assertEqual((beta["d"], beta["op"], beta["ec"]), ("2026-10-01", "2026-08-20", "register"))
        self.assertEqual([m["n"] for m in data["authorised"]], ["Betarix", "Alphazon", "Ecfallback"])
        self.assertIn("applied to 1 pending: Betarix", res.stdout)
        # the good download is remembered for outage days
        self.assertEqual(self.cache_decisions(), {"EMEA/H/C/006002": {"d": "2026-10-01", "id": "2049"},
                                 "EMEA/H/C/006004": {"d": "2026-05-01", "id": "2049"}})
        self.assertEqual(sources["euRegisterDecisions"], 2)
        self.assertFalse(sources["euRegisterFromCache"])
        self.assertEqual(sources["euRegister"]["url"], "https://ec.europa.eu/health/documents/community-register/html/reg_last.htm")

    def test_earliest_decision_wins_for_the_same_procedure(self):
        reg = os.path.join(self.dir, "reg_last.htm")
        write_register(reg, [register_record("EMEA/H/C/006002/0000", "08/10/2026", page_id="2"),
                             register_record("emea/h/c/006002/0000", "01/10/2026", page_id="1")])
        self.build(reg)
        self.assertEqual(self.cache_decisions()["EMEA/H/C/006002"], {"d": "2026-10-01", "id": "1"})

    def test_outage_page_falls_back_to_the_last_good_download(self):
        reg = os.path.join(self.dir, "reg_last.htm")
        write_register(reg, [register_record("EMEA/H/C/006002/0000", "01/10/2026")])
        self.build(reg)
        # today the Commission serves its outage page (HTTP 200, no decision records)
        with open(reg, "w", encoding="utf-8") as f:
            f.write(OUTAGE_PAGE)
        res = self.build(reg)
        data, sources = self.read()
        self.assertEqual(next(m for m in data["authorised"] if m["n"] == "Betarix")["d"], "2026-10-01")
        self.assertTrue(sources["euRegisterFromCache"])
        self.assertEqual(sources["euRegisterDecisions"], 1)
        self.assertIn("applied to 1 pending: Betarix", res.stdout)
        # the cache keeps the good decisions, the outage page never overwrites it
        self.assertEqual(list(self.cache_decisions()), ["EMEA/H/C/006002"])

    def test_missing_register_and_no_cache_leaves_pending_list_unverified(self):
        res = self.build(os.path.join(self.dir, "does-not-exist.htm"))
        data, sources = self.read()
        self.assertIn("Betarix", [m["n"] for m in data["pipeline"]])
        self.assertIn("register not available", res.stdout)
        self.assertIsNone(sources["euRegister"])


if __name__ == "__main__":
    unittest.main()
